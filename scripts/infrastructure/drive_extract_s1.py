"""
Spike S1 (USER_DRIVE_INDEXING_RFC §18, §9): do the drive parser's child-process limits hold on the indexer's
runtime? Every input goes through SubprocessContentExtractor with the given limits, each in a fresh measuring
process, and the spike prints one JSON line per file and a summary: outcome, wall time, the child's peak RSS, the
host process's peak RSS and the container cgroup's memory peak before and after the run (cgroup v2 `memory.peak`,
v1 `memory.max_usage_in_bytes`). Only the cgroup sees what RSS misses: Cloud Run's disk is memory, so the PDF windows
a split writes count against the instance. It also probes whether the kernel enforces RLIMIT_AS at all (Cloud Run gen2 must) and,
where it does, the smallest address space the parser stack starts in (the floor the limit must clear).

Run it where the indexer runs: a one-off Cloud Run Job from the indexer's image, gen2, the indexer's memory and CPU
(plan task, OWNER). Locally (macOS) RLIMIT_AS is not enforced; the other numbers are indicative only.

Inputs: the generated set (`--generate DIR`: Phase 0's ZIP member and size bombs, a sheet declaring 16384 x 1048576
cells, a long sheet, an ODS with hostile repeat counts, a 2 000-page PDF and its page windows, a 200-page PDF whose
pages all draw one 8 MiB image (every window copies it), 45 MB text and CSV
downloads, deeply nested HTML, a long DOCX, an encrypted PDF, tiny files for the child's start-up time), plus every
file in `--inputs DIR` — put the Phase 0 originals there if they still exist (66 KB DOCX, 5 KB XLSX, 8.7 KB XLS,
408 KB PDF). The kind comes from the name and the head, as the tick derives it (kind_of + sniff_kind).

The spike sets (in the plan's S1 task): EXTRACT_MEMORY_LIMIT_BYTES, EXTRACT_CPU_LIMIT_S, EXTRACT_FILE_LIMIT_BYTES
(src/adapters/subprocess_content_extractor.py), EXTRACT_WALL_TIMEOUT_S (src/domain/drive_index_params.py) and the
indexer's memory.

Usage:
    python scripts/infrastructure/drive_extract_s1.py --generate /tmp/s1 [--inputs DIR] [--memory-mib 1024]
        [--cpu-s 90] [--wall-s 120] [--file-mib 64] [--scale 1.0]
    python scripts/infrastructure/drive_extract_s1.py --one FILE --kind KIND [limits]     (one measured run)
Exit code: 0 (a spike reports, it does not judge).
"""
import argparse
import asyncio
import json
import os
import resource
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

_MIB = 1024 * 1024


def generate(target: Path, scale: float = 1.0) -> List[Path]:
    """Write the synthetic bomb set into `target`; `scale` < 1 shrinks the big ones (tests)."""
    from tests.unit.adapters.drive_extraction_fixtures import (
        encrypt_pdf, make_docx, make_inflation_bomb, make_member_bomb, make_odf, make_pdf,
    )

    target.mkdir(parents=True, exist_ok=True)
    made = [make_member_bomb(target / "phase0_member_bomb.docx", prefix="word/"),
            make_inflation_bomb(target / "phase0_inflation_bomb.docx")]
    made.append(_wide_sheet(target / "declared_16384x1048576.xlsx"))
    made.append(_long_sheet(target / "long_sheet.xlsx", rows=int(200_000 * scale)))
    made.append(make_odf(target / "hostile_repeats.ods", "ods",
                         '<office:spreadsheet><table:table table:name="S"><table:table-row '
                         'table:number-rows-repeated="1048576"><table:table-cell table:number-columns-repeated="16384">'
                         '<text:p>x</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet>'))
    made.append(make_pdf(target / "pages_2000.pdf", [f"Page {i} of a long report" for i in range(int(2000 * scale))]))
    made.append(_shared_resource_pdf(target / "shared_image_pages.pdf", pages=max(2, int(200 * scale)),
                                     image_bytes=max(1024, int(8 * _MIB * scale))))
    made.append(encrypt_pdf(make_pdf(target / "plain.pdf", ["secret"]), target / "encrypted.pdf",
                            user_password="pw"))
    made.append(_big_text(target / "download_45mb.txt", int(45 * _MIB * scale)))
    made.append(_big_csv(target / "download_45mb.csv", int(45 * _MIB * scale)))
    made.append(_nested_html(target / "nested.html", depth=int(100_000 * scale)))
    made.append(make_docx(target / "long.docx",
                          [f"Paragraph {i} of a long contract." for i in range(int(20_000 * scale))]))
    tiny = target / "tiny.txt"
    tiny.write_text("start-up time of the child", encoding="utf-8")
    made += [tiny, make_docx(target / "tiny.docx", ["start-up time with markitdown"]),
             make_pdf(target / "tiny.pdf", ["start-up time with pdfminer"])]
    return made


def _wide_sheet(path: Path) -> Path:
    import openpyxl

    book = openpyxl.Workbook()
    book.active["A1"] = "first"
    book.active["XFD1048576"] = "last"           # openpyxl writes <dimension ref="A1:XFD1048576"/>
    book.save(str(path))
    return path


def _long_sheet(path: Path, rows: int) -> Path:
    import openpyxl

    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    for i in range(rows):
        sheet.append([i, f"row {i}", i * 1.5, "2026-01-01"])
    book.save(str(path))
    return path


def _shared_resource_pdf(path: Path, pages: int, image_bytes: int) -> Path:
    """`pages` pages that all draw one image stream of `image_bytes` random bytes: small on disk, but a page window
    carries its own copy of the image, so splitting multiplies it by the number of windows."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject, StreamObject

    writer = PdfWriter()
    image = StreamObject()
    image.set_data(os.urandom(image_bytes))
    width = 1024
    image.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                  NameObject("/Width"): NumberObject(width),
                  NameObject("/Height"): NumberObject(max(1, image_bytes // (3 * width))),
                  NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                  NameObject("/BitsPerComponent"): NumberObject(8)})
    image_ref = writer._add_object(image)
    for _ in range(pages):
        page = writer.add_blank_page(612, 792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image_ref})})
        content = DecodedStreamObject()
        content.set_data(b"q 100 0 0 100 0 0 cm /Im0 Do Q")
        page[NameObject("/Contents")] = writer._add_object(content)
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def _big_text(path: Path, size: int) -> Path:
    line = "Орендна плата за квартиру 12 500,00 грн сплачується щомісяця.\n".encode("cp1251")
    with open(path, "wb") as fh:
        for _ in range(max(1, size // len(line) - 1)):
            fh.write(line)
    return path


def _big_csv(path: Path, size: int) -> Path:
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("date,amount,description\n")
        written, i = 0, 0
        while written < size - 4096:             # just under the download cap: the cap itself is tested elsewhere
            line = f"2026-01-{i % 28 + 1:02d},{i}.50,payment number {i}\n"
            fh.write(line)
            written += len(line)
            i += 1
    return path


def _nested_html(path: Path, depth: int) -> Path:
    path.write_text("<html><body>" + "<div>" * depth + "deep" + "</div>" * depth + "</body></html>",
                    encoding="utf-8")
    return path


def kind_for(path: Path) -> str:
    from src.domain.drive_index_policy import kind_of
    from src.domain.drive_sniff import SNIFF_HEAD_BYTES, sniff_kind

    with open(path, "rb") as fh:
        head = fh.read(SNIFF_HEAD_BYTES)
    return sniff_kind(head, kind_of(path.name, "", False))


def rlimit_probe(limit_mib: int = 256) -> Dict[str, Any]:
    """Does this kernel enforce RLIMIT_AS? A child sets it and allocates twice as much."""
    code = (f"import resource; resource.setrlimit(resource.RLIMIT_AS, ({limit_mib * _MIB}, {limit_mib * _MIB}))\n"
            f"try:\n    b = bytearray({2 * limit_mib * _MIB}); print('allocated')\n"
            f"except MemoryError:\n    print('refused')")
    proc = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=60)
    outcome = proc.stdout.strip() or (proc.stderr.strip().splitlines() or [""])[-1]
    return {"platform": sys.platform, "rlimit_as_honoured": outcome == "refused", "probe": outcome}


async def _measure_one(path: str, kind: str, args: argparse.Namespace) -> Dict[str, Any]:
    from src.adapters.subprocess_content_extractor import SubprocessContentExtractor
    from src.domain.drive_extraction import ExtractionRequest, PdfSplitFailed

    extractor = SubprocessContentExtractor(memory_limit_bytes=args.memory_mib * _MIB, cpu_limit_s=args.cpu_s,
                                           wall_timeout_s=args.wall_s, file_limit_bytes=args.file_mib * _MIB)
    started = time.monotonic()
    row: Dict[str, Any] = {"file": os.path.basename(path), "kind": kind, "bytes": os.path.getsize(path)}
    if args.split:
        with tempfile.TemporaryDirectory() as out_dir:
            try:   # as the reader does: the page count, then one window at a time, removed once cut
                pages_total = await extractor.pdf_page_count(path)
                windows = 0
                for first_page in range(0, min(pages_total, 200), 10):
                    window = os.path.join(out_dir, "window.pdf")
                    await extractor.cut_pdf_window(path, first_page, 10, window)
                    os.remove(window)
                    windows += 1
                row.update(status="split", windows=windows, pages_total=pages_total, truncated=pages_total > 200)
            except PdfSplitFailed as exc:
                row.update(status=exc.status.value, detail=exc.detail)
    else:
        outcome = await extractor.extract(ExtractionRequest(path=path, name=os.path.basename(path), kind=kind,
                                                            mime_type="", user_id="s1", account_id="s1"))
        row.update(status=outcome.status.value, detail=outcome.detail[:120], text_chars=len(outcome.text),
                   pages_total=outcome.pages_total, truncated=outcome.truncated)
    row["seconds"] = round(time.monotonic() - started, 2)
    row["child_peak_mib"] = _peak_mib(resource.RUSAGE_CHILDREN)
    row["host_peak_mib"] = _peak_mib(resource.RUSAGE_SELF)
    return row


def _peak_mib(who: int) -> float:
    peak = resource.getrusage(who).ru_maxrss
    return round(peak / (_MIB if sys.platform == "darwin" else 1024), 1)   # bytes on macOS, KiB on Linux


def cgroup_memory_peak_bytes(root: Path = Path("/sys/fs/cgroup")) -> Optional[int]:
    """The container's memory high-water mark (page cache and the in-memory disk included): cgroup v2, else v1,
    else None (not Linux, or no cgroup memory controller visible)."""
    for name in ("memory.peak", "memory/memory.max_usage_in_bytes"):
        try:
            return int((root / name).read_text().strip())
        except (OSError, ValueError):
            continue
    return None


def _mib_or_none(value: Optional[int]) -> Optional[float]:
    return None if value is None else round(value / _MIB, 1)


def _run_measured(path: Path, kind: str, args: argparse.Namespace, *, split: bool = False) -> Dict[str, Any]:
    """One fresh process per file, so RUSAGE_CHILDREN is that file's child alone. The cgroup peak is read around
    it: a rise is this run's (the peak never falls, so an unchanged value means the run stayed below it)."""
    command = [sys.executable, str(Path(__file__).resolve()), "--one", str(path), "--kind", kind,
               "--memory-mib", str(args.memory_mib), "--cpu-s", str(args.cpu_s), "--wall-s", str(args.wall_s),
               "--file-mib", str(args.file_mib)] + (["--split"] if split else [])
    before = cgroup_memory_peak_bytes()
    proc = subprocess.run(command, capture_output=True, text=True, timeout=args.wall_s + 120)
    after = cgroup_memory_peak_bytes()
    lines = [line for line in proc.stdout.splitlines() if line.startswith("{")]
    if lines:
        row = dict(json.loads(lines[-1]))
    else:
        row = {"file": path.name, "kind": kind, "status": "script_failed", "detail": proc.stderr[-300:]}
    row["cgroup_peak_before_mib"] = _mib_or_none(before)
    row["cgroup_peak_after_mib"] = _mib_or_none(after)
    return row


def run(paths: List[Path], args: argparse.Namespace) -> List[Dict[str, Any]]:
    rows = []
    for path in paths:
        kind = kind_for(path)
        rows.append(_run_measured(path, kind, args))
        print(json.dumps(rows[-1], ensure_ascii=False), flush=True)
        if kind == "pdf" and path.name.startswith(("pages_", "shared_image_")):
            rows.append(_run_measured(path, kind, args, split=True))
            print(json.dumps(rows[-1], ensure_ascii=False), flush=True)
    return rows


def address_space_floor(path: Path, kind: str, args: argparse.Namespace, low: int = 64,
                        high: int = 4096) -> Dict[str, Any]:
    """The smallest RLIMIT_AS (MiB, to 32 MiB) under which this small file still reads: the parser stack's own
    address space, which the limit must clear by a wide margin. Meaningful only where RLIMIT_AS is honoured."""
    def reads(mib: int) -> bool:
        trial = argparse.Namespace(**{**vars(args), "memory_mib": mib})
        return bool(_run_measured(path, kind, trial)["status"] == "text")

    if not reads(high):
        return {"file": path.name, "floor_mib": None}
    while high - low > 32:
        middle = (low + high) // 2
        if reads(middle):
            high = middle
        else:
            low = middle
    return {"file": path.name, "floor_mib": high}


def summary(rows: List[Dict[str, Any]], probe: Dict[str, Any], args: argparse.Namespace) -> str:
    head = (f"S1 on {probe['platform']}: RLIMIT_AS honoured={probe['rlimit_as_honoured']}; limits: "
            f"AS {args.memory_mib} MiB, CPU {args.cpu_s} s, wall {args.wall_s} s, file {args.file_mib} MiB")
    lines = [head, f"{'file':38} {'kind':9} {'status':12} {'sec':>7} {'child MiB':>10} {'host MiB':>9} "
                   f"{'cgroup peak MiB':>17}"]
    for row in rows:
        lines.append(f"{row['file'][:38]:38} {row['kind']:9} {row['status']:12} {row.get('seconds', 0):>7} "
                     f"{row.get('child_peak_mib', 0):>10} {row.get('host_peak_mib', 0):>9} "
                     f"{_peak_change(row):>17}")
    return "\n".join(lines)


def _peak_change(row: Dict[str, Any]) -> str:
    before, after = row.get("cgroup_peak_before_mib"), row.get("cgroup_peak_after_mib")
    if after is None:
        return "n/a"
    return f"{before}->{after}" if before != after else f"{after} (=)"


def _parse(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--generate", type=Path, help="write the synthetic set here and run it")
    parser.add_argument("--inputs", type=Path, help="also run every file in this directory")
    parser.add_argument("--one", help=argparse.SUPPRESS)
    parser.add_argument("--kind", help=argparse.SUPPRESS)
    parser.add_argument("--split", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--memory-mib", type=int, default=1024)
    parser.add_argument("--cpu-s", type=int, default=90)
    parser.add_argument("--wall-s", type=float, default=120)
    parser.add_argument("--file-mib", type=int, default=64)
    parser.add_argument("--scale", type=float, default=1.0)
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = _parse(argv)
    if args.one:
        print(json.dumps(asyncio.run(_measure_one(args.one, args.kind, args)), ensure_ascii=False))
        return 0
    paths: List[Path] = []
    if args.generate:
        paths += generate(args.generate, args.scale)
    if args.inputs:
        paths += sorted(p for p in args.inputs.iterdir() if p.is_file())
    probe = rlimit_probe()
    print(json.dumps(probe), flush=True)
    rows = run(paths, args)
    if probe["rlimit_as_honoured"]:
        for path in (p for p in paths if p.name in ("tiny.docx", "tiny.pdf")):
            print(json.dumps(address_space_floor(path, kind_for(path), args)), flush=True)
    print(summary(rows, probe, args))
    return 0


if __name__ == "__main__":
    sys.exit(main())
