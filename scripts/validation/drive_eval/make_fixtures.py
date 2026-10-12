"""
Generate the eval's content fixtures (USER_DRIVE_INDEXING_RFC §17.8, §24.11) into tests/fixtures/drive_eval/files/.

    venv/bin/python -m scripts.validation.drive_eval.make_fixtures            # write every file
    venv/bin/python -m scripts.validation.drive_eval.make_fixtures --only scan_0012.pdf

The output is committed, not generated at run time: Chromium's rendering (fonts, anti-aliasing, PDF object order)
differs between machines and versions, and the recorded bars pin the bytes they were measured on (`files_sha256` in
bars.json). Re-running this script is a deliberate fixture change that is followed by `--record`. The text formats are
byte-deterministic; the rendered ones (PDFs printed by Chromium, scans, screenshots, photos) are not.

Every person, organisation, number and account in the files is fictional (fixture_texts.py).
Needs Playwright's Chromium (`venv/bin/playwright install chromium`), python-docx, openpyxl, python-pptx, pypdf, Pillow.
"""
from __future__ import annotations

import argparse
import csv
import io
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Sequence, Tuple
from xml.sax.saxutils import escape

from . import fixture_texts as T

OUT = Path(__file__).resolve().parents[3] / "tests" / "fixtures" / "drive_eval" / "files"
FIXED_TIME = (2026, 1, 15, 9, 0, 0)            # zip entry times: the text formats come out byte-identical


# -- deterministic text formats ------------------------------------------------------------------------------------

def _zip(entries: Sequence[Tuple[str, bytes]], stored_first: bool = False) -> bytes:
    """A zip with fixed entry times; ODF wants `mimetype` first and uncompressed."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i, (name, data) in enumerate(entries):
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.compress_type = zipfile.ZIP_STORED if stored_first and i == 0 else zipfile.ZIP_DEFLATED
            zf.writestr(info, data)
    return buf.getvalue()


_ODF_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
           'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
           'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
           'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
           'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
           'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"')


def _odf(mimetype: str, body: str) -> bytes:
    content = (f'<?xml version="1.0" encoding="UTF-8"?>\n<office:document-content {_ODF_NS} office:version="1.3">'
               f"<office:body>{body}</office:body></office:document-content>")
    manifest = (f'<?xml version="1.0" encoding="UTF-8"?>\n<manifest:manifest {_ODF_NS} manifest:version="1.3">'
                f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mimetype}"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                "</manifest:manifest>")
    return _zip([("mimetype", mimetype.encode()), ("content.xml", content.encode("utf-8")),
                 ("META-INF/manifest.xml", manifest.encode("utf-8"))], stored_first=True)


def odt_bytes(paragraphs: Sequence[str]) -> bytes:
    body = "".join(f"<text:p>{escape(p)}</text:p>" for p in paragraphs)
    return _odf("application/vnd.oasis.opendocument.text", f"<office:text>{body}</office:text>")


def ods_bytes(rows: Sequence[Sequence[object]], sheet: str = "Budget") -> bytes:
    def cell(value: object) -> str:
        if isinstance(value, (int, float)):
            return (f'<table:table-cell office:value-type="float" office:value="{value}">'
                    f"<text:p>{value}</text:p></table:table-cell>")
        return f'<table:table-cell office:value-type="string"><text:p>{escape(str(value))}</text:p></table:table-cell>'
    body = "".join("<table:table-row>" + "".join(cell(v) for v in row) + "</table:table-row>" for row in rows)
    return _odf("application/vnd.oasis.opendocument.spreadsheet",
                f'<office:spreadsheet><table:table table:name="{escape(sheet)}">{body}</table:table>'
                "</office:spreadsheet>")


def odp_bytes(slides: Sequence[Tuple[str, str]]) -> bytes:
    pages = "".join(
        f'<draw:page draw:name="page{i + 1}"><draw:frame svg:width="20cm" svg:height="3cm"><draw:text-box>'
        f"<text:p>{escape(title)}</text:p><text:p>{escape(body)}</text:p></draw:text-box></draw:frame></draw:page>"
        for i, (title, body) in enumerate(slides))
    return _odf("application/vnd.oasis.opendocument.presentation", f"<office:presentation>{pages}</office:presentation>")


def rtf_bytes(paragraphs: Sequence[str]) -> bytes:
    """Non-ASCII as \\uN? escapes, as word processors write Cyrillic RTF."""
    def enc(text: str) -> str:
        out = []
        for ch in text:
            code = ord(ch)
            if ch in "\\{}":
                out.append("\\" + ch)
            elif code < 128:
                out.append(ch)
            else:
                out.append(f"\\u{code if code < 32768 else code - 65536}?")
        return "".join(out)
    body = "".join(f"{enc(p)}\\par\n" for p in paragraphs)
    return ("{\\rtf1\\ansi\\ansicpg1251\\deff0{\\fonttbl{\\f0 Times New Roman;}}\\f0\\fs24\n" + body + "}").encode("ascii")


def csv_bytes(header: Sequence[str], rows: Sequence[Sequence[str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def text_bytes(text: str) -> bytes:
    return text.encode("utf-8")


# -- office formats (python-docx, openpyxl, python-pptx) -------------------------------------------------------------

def _fixed_core(props: object) -> None:
    for name in ("created", "modified", "last_printed"):
        if hasattr(props, name):
            setattr(props, name, datetime(*FIXED_TIME))


def docx_bytes(paragraphs: Sequence[str], table: Sequence[Sequence[str]] = ()) -> bytes:
    from docx import Document
    doc = Document()
    _fixed_core(doc.core_properties)
    doc.add_heading(paragraphs[0], level=1)
    for p in paragraphs[1:]:
        doc.add_paragraph(p)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def xlsx_bytes(sheet: str, rows: Sequence[Sequence[object]]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(list(row))
    _fixed_core(wb.properties)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pptx_bytes(slides: Sequence[Tuple[str, str]]) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    _fixed_core(prs.core_properties)
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# -- rendered formats (Chromium through Playwright) -----------------------------------------------------------------

_PAGE_CSS = ("body{font-family:Arial,'Helvetica Neue','DejaVu Sans',sans-serif;font-size:12.5px;line-height:1.45;"
             "margin:48px 56px;color:#111} h1{font-size:20px} h2{font-size:15px;margin-top:18px} "
             "table{border-collapse:collapse;width:100%} td,th{border:1px solid #444;padding:4px 6px;text-align:left}"
             " .cols{column-count:2;column-gap:28px;text-align:justify}")


class Renderer:
    """One Chromium for the whole run."""

    def __enter__(self) -> "Renderer":
        from playwright.sync_api import sync_playwright
        self._pw = sync_playwright().start()
        self._browser = self._pw.chromium.launch()
        return self

    def __exit__(self, *exc: object) -> None:
        self._browser.close()
        self._pw.stop()

    def pdf(self, html: str) -> bytes:
        page = self._browser.new_page()
        page.set_content(f"<html><head><meta charset='utf-8'><style>{_PAGE_CSS}</style></head><body>{html}</body></html>")
        data = page.pdf(format="A4", print_background=True)
        page.close()
        return bytes(data)

    def pixels(self, html: str, width: int, height: int, fmt: str = "png", quality: int = 72) -> bytes:
        page = self._browser.new_page(viewport={"width": width, "height": height})
        page.set_content(f"<html><head><meta charset='utf-8'></head><body style='margin:0'>{html}</body></html>")
        page.wait_for_timeout(100)
        kwargs: Dict[str, object] = {"type": "jpeg", "quality": quality} if fmt == "jpeg" else {"type": "png"}
        data = page.screenshot(full_page=False, **kwargs)  # type: ignore[arg-type]
        page.close()
        return bytes(data)


def _lines_html(lines: Sequence[Tuple[str, str]]) -> str:
    tags = {"h": "<div style='font-size:1.35em;font-weight:700;margin-bottom:.6em'>{}</div>",
            "b": "<div style='font-weight:700;font-size:1.15em;margin:.5em 0'>{}</div>", "p": "<div>{}</div>"}
    return "".join(tags[kind].format(escape(text)) for kind, text in lines)


def paper_scan_html(lines: Sequence[Tuple[str, str]], mono: bool = False) -> str:
    """A sheet photographed or scanned slightly askew: greyish paper, soft focus, no text layer."""
    font = "'Courier New',monospace" if mono else "Arial,'DejaVu Sans',sans-serif"
    return (f"<div style='width:100%;height:100%;background:#9a9a96;display:flex;align-items:center;"
            f"justify-content:center'><div style='width:78%;padding:46px 40px;background:#f1efe8;"
            f"transform:rotate(-1.2deg);filter:blur(.45px) contrast(1.05);font-family:{font};font-size:23px;"
            f"line-height:1.6;color:#222;box-shadow:0 6px 18px rgba(0,0,0,.35)'>{_lines_html(lines)}</div></div>")


def phone_screen_html(lines: Sequence[Tuple[str, str]], accent: str) -> str:
    return (f"<div style='width:100%;height:100%;background:#f4f5f7;font-family:-apple-system,Arial,sans-serif'>"
            f"<div style='height:44px;background:{accent}'></div>"
            f"<div style='margin:22px;padding:26px;background:#fff;border-radius:18px;font-size:22px;line-height:1.7;"
            f"color:#1d1d1f;box-shadow:0 2px 8px rgba(0,0,0,.08)'>{_lines_html(lines)}</div></div>")


def chat_bubble_html(lines: Sequence[Tuple[str, str]]) -> str:
    return ("<div style='width:100%;height:100%;background:#e5ddd5;font-family:Arial,sans-serif;padding-top:30px'>"
            "<div style='margin:0 70px 0 18px;padding:18px 22px;background:#fff;border-radius:14px;font-size:21px;"
            f"line-height:1.6'>{_lines_html(lines)}<div style='text-align:right;color:#888;font-size:15px'>10:02</div>"
            "</div></div>")


def scene_html(sky: str, ground: str, shapes: str) -> str:
    """A flat, wordless picture: the triage model should call it a photograph, never text."""
    return (f"<div style='position:relative;width:100%;height:100%;background:linear-gradient({sky})'>"
            f"<div style='position:absolute;bottom:0;width:100%;height:38%;background:{ground}'></div>{shapes}</div>")


def image_pdf(jpeg: bytes) -> bytes:
    """A scan: the page is one picture, so the PDF has no text layer at all."""
    from PIL import Image
    buf = io.BytesIO()
    Image.open(io.BytesIO(jpeg)).convert("RGB").save(buf, "PDF", resolution=150.0)
    return buf.getvalue()


def strip_unicode_maps(pdf: bytes) -> bytes:
    """A PDF that looks right and extracts as (cid:N): no /ToUnicode, and the embedded fonts' `cmap` table renamed.
    Renderers draw Identity-H CID fonts through glyph ids, so the page is intact; text extraction has no way back to
    Unicode (§9 router signal "garbage text")."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import NameObject
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf)))
    for page in writer.pages:
        fonts = page["/Resources"]["/Font"]
        for key in fonts:
            font = fonts[key].get_object()
            if "/ToUnicode" in font:
                del font[NameObject("/ToUnicode")]
            for descendant in font.get("/DescendantFonts", []):
                program = descendant.get_object()["/FontDescriptor"].get_object().get("/FontFile2")
                if program is None:
                    continue
                stream = program.get_object()
                data = bytearray(stream.get_data())
                for i in range(int.from_bytes(data[4:6], "big")):
                    tag = 12 + 16 * i
                    if data[tag:tag + 4] == b"cmap":
                        data[tag:tag + 4] = b"xmap"
                stream.set_data(bytes(data))
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# -- the fixture set ----------------------------------------------------------------------------------------------

Builder = Callable[[Renderer], bytes]


def _dogovor(version: str) -> Builder:
    return lambda r: docx_bytes(T.dogovor_paragraphs(version))


def _scan(lines: Sequence[Tuple[str, str]], mono: bool = False) -> Builder:
    return lambda r: image_pdf(r.pixels(paper_scan_html(lines, mono), 1240, 1754, "jpeg", 70))


FIXTURES: Dict[str, Builder] = {
    # documents with a text layer
    "dogovor_red1.docx": _dogovor("1"),
    "dogovor_red2.docx": _dogovor("2"),
    "dogovor_red3.docx": _dogovor("3"),
    "contrato_alquiler_2026.pdf": lambda r: r.pdf(T.CONTRATO_HTML),
    "lease_porto.pdf": lambda r: r.pdf(T.LEASE_PORTO_HTML),                 # two columns
    "super_important_agreement.pdf": lambda r: r.pdf(T.NDA_HTML),
    "resume.pdf": lambda r: r.pdf(T.RESUME_HTML),
    "coeur.pdf": lambda r: r.pdf(T.COEUR_HTML),
    "tax_summary_2025.pdf": lambda r: r.pdf(T.TAX_SUMMARY_HTML),           # table-heavy
    "schet_faktura_1847.pdf": lambda r: strip_unicode_maps(r.pdf(T.SCHET_FAKTURA_HTML)),   # broken-font Cyrillic
    "quarterly_planning_review.docx": lambda r: docx_bytes(T.PLANNING_REVIEW_INTRO, T.PLANNING_REVIEW_TABLE),
    "quarterly_review.pptx": lambda r: pptx_bytes(T.QUARTERLY_REVIEW_SLIDES),
    "roadmap_2027.odp": lambda r: odp_bytes(T.ROADMAP_SLIDES),
    "akt_52.odt": lambda r: odt_bytes(T.AKT_PARAGRAPHS),
    "budget_2026.ods": lambda r: ods_bytes([("Category", "EUR"), *T.BUDGET_2026]),
    "invoice_2026_03.xlsx": lambda r: xlsx_bytes("Invoice", [
        ("Invoice", "INV-2026-0317"), ("Date", "2026-03-31"), ("Customer", "Sample Workshop SL"),
        ("Item", "Quantity", "Unit price", "Amount"),
        *[(item, qty, price, round(qty * price, 2)) for item, qty, price in T.INVOICE_2026_03_ROWS],
        ("Total", "", "", T.INVOICE_2026_03_TOTAL)]),
    "sales_regions_2026.xlsx": lambda r: xlsx_bytes("Sales", [("Region", "Q1", "Q2", "Q3", "Q4"), *T.SALES_REGIONS]),
    "expenses_2025.csv": lambda r: csv_bytes(("date", "category", "description", "amount_eur"), T.EXPENSES_CSV_ROWS),
    "notes_2.md": lambda r: text_bytes(T.NOTES_2_MD),
    "idei_proektu.md": lambda r: text_bytes(T.IDEI_MD),
    "plan_viaje_roma.md": lambda r: text_bytes(T.PLAN_VIAJE_MD),
    "fiestas_del_ano.md": lambda r: text_bytes(T.FIESTAS_MD),
    "weekly_digest.html": lambda r: text_bytes(T.WEEKLY_DIGEST_HTML),
    "rent_november.eml": lambda r: text_bytes(T.RENT_EML),
    "meeting_2026_09_14.vtt": lambda r: text_bytes(T.MEETING_VTT),
    "ustav_kluba.rtf": lambda r: rtf_bytes(T.USTAV_PARAGRAPHS),
    "site_export_package.json": lambda r: text_bytes(T.SITE_EXPORT_PACKAGE_JSON),
    # no text layer: scans, a receipt photo, screenshots (Slice 4, the reader)
    "scan_0012.pdf": _scan(T.SCAN_ELECTRICITY_UK),
    "scan_0013.pdf": _scan(T.SCAN_HARDWARE_ES, mono=True),
    "img_4912.jpg": lambda r: r.pixels(paper_scan_html(T.RECEIPT_CAFE_RU, mono=True), 720, 960, "jpeg", 70),
    "screenshot_bank_en.png": lambda r: r.pixels(phone_screen_html(T.SCREEN_BANK_EN, "#0b6e4f"), 430, 720),
    "screenshot_parcel_uk.png": lambda r: r.pixels(phone_screen_html(T.SCREEN_PARCEL_UK, "#d7263d"), 430, 720),
    "screenshot_train_es.png": lambda r: r.pixels(phone_screen_html(T.SCREEN_TRAIN_ES, "#6a2c91"), 430, 720),
    "wa_dentist_ru.jpg": lambda r: r.pixels(chat_bubble_html(T.SCREEN_DENTIST_RU), 540, 520, "jpeg", 72),
    # wordless pictures: triage captions them as photographs (§24.5)
    "photo_rome.jpg": lambda r: r.pixels(scene_html(
        "#f6c27a,#f9e1b5", "#b9875a",
        "<div style='position:absolute;left:18%;bottom:30%;width:64%;height:34%;background:#d9b48a;"
        "border-radius:50% 50% 0 0'></div><div style='position:absolute;left:12%;top:12%;width:70px;height:70px;"
        "border-radius:50%;background:#fff3c4'></div>"), 640, 480, "jpeg", 70),
    "photo_porto.jpg": lambda r: r.pixels(scene_html(
        "#7fb3d5,#d6eaf8", "#2e86c1",
        "<div style='position:absolute;left:8%;bottom:34%;width:22%;height:30%;background:#e59866'></div>"
        "<div style='position:absolute;left:34%;bottom:34%;width:18%;height:38%;background:#f5cba7'></div>"
        "<div style='position:absolute;left:56%;bottom:34%;width:26%;height:26%;background:#dc7633'></div>"),
        640, 480, "jpeg", 70),
    "photo_party.jpg": lambda r: r.pixels(scene_html(
        "#2c3e50,#4a235a", "#1b2631",
        "".join(f"<div style='position:absolute;left:{10 + 16 * i}%;top:{20 + (i % 2) * 14}%;width:46px;"
                f"height:46px;border-radius:50%;background:hsl({i * 60},80%,60%)'></div>" for i in range(5))),
        640, 480, "jpeg", 70),
}


def write_fixtures(out: Path, only: Sequence[str] = ()) -> List[Path]:
    unknown = sorted(set(only) - set(FIXTURES))
    if unknown:
        raise SystemExit(f"unknown fixture(s): {', '.join(unknown)}")
    out.mkdir(parents=True, exist_ok=True)
    written: List[Path] = []
    with Renderer() as renderer:
        for name, build in FIXTURES.items():
            if only and name not in only:
                continue
            path = out / name
            path.write_bytes(build(renderer))
            written.append(path)
    return written


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m scripts.validation.drive_eval.make_fixtures")
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--only", action="append", default=[], help="one file name; repeatable")
    args = parser.parse_args(argv)
    for path in write_fixtures(args.out, args.only):
        sys.stdout.write(f"{path.stat().st_size:>8}  {path.name}\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
