"""
Per-kind text readers of the drive parser (docs/10_rfcs/USER_DRIVE_INDEXING_RFC.md §5.1, §8.3, §9, §24.3).

They run only inside the sandboxed child process (`drive_parser.child`), never in the indexer process: every
parser here may amplify a small file into gigabytes (§9, Phase 0). Parser libraries are imported inside the
functions, so they load after the child has applied its limits.

Every admitted kind has a named reader (§5.1): pdfminer.six for PDF, called directly (markitdown 0.1.8's PdfConverter
runs pdfplumber first and falls back to pdfminer on any exception: two engines, one of them unpinned by intent);
markitdown's allowlisted converters (Phase 0, `src/utils/document_hardening.build_markitdown`) for DOCX, PPTX, EPUB,
MSG and HTML; charset detection for the
text kinds, VTT/SRT reduced to cue text; `striprtf` for RTF; the stdlib `email` parser for EML; ODF `content.xml`
through `defusedxml`; XLSX/XLS/CSV/ODS as Markdown tables capped at MAX_SHEET_ROWS rows per sheet (the chunker
repeats the header row, §8.3). Password-protected PDF/Office/EPUB/ODF → `encrypted`. An EPUB that is a plain tar
(an Apple Books `.epub.pkgf` package) is repacked as an EPUB ZIP in memory first (`epub_package`).

Parser@2 (§25.3, D23, review thresholds): HTML pages, and Markdown or text files that are HTML (`html_text.text_is_html`:
a document marker or tags over 30 % outside code), go through `html_text.html_page_text` — the content landmark,
no top-level boilerplate or in-page link blocks, metadata lines at the head — and a page with a <script> and under
500 visible characters comes back `script_rendered`. A text kind whose HTML path fails keeps its decoded text.

Public interface:
    MARKITDOWN_CONVERTERS
    read_content(path, kind) -> ExtractionOutcome          # `extracted_by` is stamped by the parent
    pdf_page_count(path) -> int                                   # raises PdfSplitFailed
    cut_pdf_window(path, first_page, page_count, out_path) -> int   # pages written; raises PdfSplitFailed
    decode_text(data, cut) -> str
    subtitle_cues(text) -> str
    markdown_table(rows) -> Tuple[str, bool]
"""
from __future__ import annotations

import codecs
import csv
import datetime as dt
import html
import io
import os
import re
import zipfile
from contextlib import contextmanager
from dataclasses import replace
from typing import Any, Callable, Dict, FrozenSet, Iterable, List, Mapping, NamedTuple, Optional, Sequence, Tuple
from typing import BinaryIO, Iterator, Union

from src.adapters.drive_parser.capture import pdf_capture_signals
from src.adapters.drive_parser.child import is_memory_failure
from src.adapters.drive_parser.epub_package import PackageError, epub_zip_from_tar, is_tar_package
from src.adapters.drive_parser.html_text import html_page_text, text_is_html
from src.domain.drive_capture import CaptureSignals
from src.domain.drive_extraction import ExtractionOutcome, ExtractionStatus, PdfSplitFailed, failure_label
from src.domain.drive_index_params import MAX_SHEET_ROWS, MAX_TEXT_CHARS
from src.utils.document_hardening import ArchiveTooLargeError, check_zip_container

# kind -> markitdown converter class (all registered by build_markitdown; a test keeps the two in step). Called
# directly with the sniffed kind: no magika guess, no fall-through to another converter.
MARKITDOWN_CONVERTERS: Mapping[str, Tuple[str, str, str]] = {
    "docx": ("DocxConverter", ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pptx": ("PptxConverter", ".pptx",
             "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "epub": ("EpubConverter", ".epub", "application/epub+zip"),
    "msg": ("OutlookMsgConverter", ".msg", "application/vnd.ms-outlook"),
}
_ZIP_KINDS: FrozenSet[str] = frozenset({"docx", "xlsx", "pptx", "epub", "odt", "ods", "odp"})
_OOXML_KINDS: FrozenSet[str] = frozenset({"docx", "xlsx", "pptx"})
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_FONT_SUFFIXES = (".ttf", ".otf", ".woff", ".woff2")   # EPUB font obfuscation is not DRM
_MAX_TABLE_COLUMNS = 1024    # bounds a hostile repeat count or dimension; real sheets stay far below
_CSV_SNIFF_CHARS = 4096
# Legacy code pages of the languages the bot serves (uk, ru, en, es, fr), in tie-break order.
_PREFERRED_LEGACY = ("cp1251", "cp1252", "cp1250", "koi8_u", "koi8_r", "iso8859_15", "latin_1")
_CHAOS_MARGIN = 0.1
_PLAUSIBILITY_SAMPLE_BYTES = 4096
_ODF = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
}
_Read = Tuple[str, int, bool]    # text, pages_total, truncated
# The kinds whose text may be a whole HTML page saved under another extension (§25.3, parser@2); never JSON.
_HTML_CAPABLE_TEXT_KINDS: FrozenSet[str] = frozenset({"text", "markdown"})


class _PageRead(NamedTuple):
    """What an HTML-converting reader returns: a _Read plus the script-rendered flag (§25.3)."""
    text: str
    pages_total: int
    truncated: bool
    script_rendered: bool


class _Encrypted(Exception):
    """The file needs a password we do not have."""


class _WrongContainer(ValueError):
    """The bytes are not the container their kind needs (an OLE2 file under a ZIP kind)."""


# -- entry points ---------------------------------------------------------------------------------------------------

def read_content(path: str, kind: str) -> ExtractionOutcome:
    """Text of one file. A per-file failure is an outcome; a memory failure propagates as MemoryError so the child
    exits with its memory code (§8.2: `memory` is per-file transient, never `error`)."""
    reader = _READERS.get(kind)
    if reader is None:
        return ExtractionOutcome(status=ExtractionStatus.UNSUPPORTED, detail=f"no reader for kind {kind}")
    try:
        if kind == "epub" and is_tar_package(path):
            result: Union[_Read, _PageRead] = _read_epub_package(path)
        else:
            if kind in _ZIP_KINDS:
                _check_container(path, kind)
            result = reader(path, kind)
    except ArchiveTooLargeError as exc:
        return ExtractionOutcome(status=ExtractionStatus.TOO_LARGE, detail=str(exc))
    except _Encrypted as exc:
        return ExtractionOutcome(status=ExtractionStatus.ENCRYPTED, detail=str(exc))
    except Exception as exc:  # any parser failure on hostile input is this file's outcome, never the tick's
        if is_memory_failure(exc):
            raise MemoryError(str(exc)) from exc
        status = ExtractionStatus.UNSUPPORTED if _deterministic(exc) else ExtractionStatus.ERROR
        return ExtractionOutcome(status=status, detail=failure_label(exc))   # never the message (R3 F3-5)
    text, pages_total, truncated = result[0], result[1], result[2]
    outcome = _finish(text, pages_total, truncated,
                      script_rendered=isinstance(result, _PageRead) and result.script_rendered)
    if kind == "pdf":
        outcome = replace(outcome, capture_signals=_capture_signals(path))
    return outcome


def _capture_signals(path: str) -> Optional[CaptureSignals]:
    """Parser@3's capture signals (§26.3). A failure to read them is no signals (the model decides, when one reads
    the file), never the file's outcome: the text was read. A memory failure still ends the child."""
    try:
        return pdf_capture_signals(path)
    except Exception as exc:  # hostile input: whatever pypdf raises on it is this read's, not the file's
        if is_memory_failure(exc):
            raise MemoryError(str(exc)) from exc
        return None


def pdf_page_count(path: str) -> int:
    """Pages of a PDF the reader may cut (§24.3), counted from the page tree like `_read_pdf` (an encrypted file's
    /Count can claim anything). An owner-password-only PDF counts; a user password is `encrypted`."""
    with _pdf_failures():
        return _page_count(_open_pdf(path))


def cut_pdf_window(path: str, first_page: int, page_count: int, out_path: str) -> int:
    """Write pages [first_page, first_page + page_count) of the PDF (0-based, clamped to its end) as one window PDF
    at `out_path`; returns the pages written. One window per call: the reader cuts on demand, so at most
    MODEL_IN_FLIGHT windows exist on disk at once (§24.3) — a window can carry a resource every page shares."""
    with _pdf_failures():
        from pypdf import PdfWriter

        reader = _open_pdf(path)
        total = _page_count(reader)                 # the page tree, never /Count (arbiter R1-1)
        if first_page < 0 or page_count < 1 or first_page >= total:
            raise PdfSplitFailed(ExtractionStatus.ERROR, f"pages {first_page}+{page_count} outside 0..{total}")
        last = min(first_page + page_count, total)
        pages = reader.flattened_pages or []           # the walked tree: `reader.pages[i]` answers from /Count too
        writer = PdfWriter()
        for index in range(first_page, last):
            writer.add_page(pages[index])
        with open(out_path, "wb") as fh:
            writer.write(fh)
        return last - first_page


def _open_pdf(path: str) -> Any:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=False)
    if reader.is_encrypted and not reader.decrypt(""):
        raise PdfSplitFailed(ExtractionStatus.ENCRYPTED, "user password")
    return reader


@contextmanager
def _pdf_failures() -> Iterator[None]:
    """A failure of the page-window ops as PdfSplitFailed with its outcome status; a memory failure propagates as
    MemoryError so the child exits with its memory code."""
    try:
        yield
    except PdfSplitFailed:
        raise
    except Exception as exc:  # any parser failure on hostile input is this file's outcome, never the tick's
        if is_memory_failure(exc):
            raise MemoryError(str(exc)) from exc
        status = ExtractionStatus.UNSUPPORTED if _deterministic(exc) else ExtractionStatus.ERROR
        raise PdfSplitFailed(status, failure_label(exc)) from exc   # never the message (R3 F3-5)


# -- shared helpers -------------------------------------------------------------------------------------------------

def decode_text(data: bytes, cut: bool = False) -> str:
    """BOM first, then strict UTF-8, then charset-normalizer's best guess (§5.1). `cut`: `data` is a prefix of the
    file, so a character split at its end is dropped instead of replaced."""
    for bom, encoding in ((codecs.BOM_UTF8, "utf-8"), (codecs.BOM_UTF32_LE, "utf-32-le"),
                          (codecs.BOM_UTF32_BE, "utf-32-be"), (codecs.BOM_UTF16_LE, "utf-16-le"),
                          (codecs.BOM_UTF16_BE, "utf-16-be")):
        if data.startswith(bom):
            return _decode(data[len(bom):], encoding, cut, "replace")
    try:
        return _decode(data, "utf-8", cut, "strict")
    except UnicodeDecodeError:
        return _decode(data, _guess_encoding(data), cut, "replace")


def subtitle_cues(text: str) -> str:
    """VTT/SRT reduced to cue text (§8.1 step 6): header, NOTE/STYLE/REGION blocks, cue ids, timing lines and blank
    separators go; `<v Name>` becomes `Name: `; other tags go; a rolled-up repeat of the previous line goes."""
    lines = [line.strip() for line in text.splitlines()]
    out: List[str] = []
    skipping = False
    for index, line in enumerate(lines):
        if not line:
            skipping = False
            continue
        if skipping or "-->" in line:
            continue
        if line.startswith("WEBVTT") or re.match(r"(NOTE|STYLE|REGION)(\s|$)", line):
            skipping = True
            continue
        if index + 1 < len(lines) and "-->" in lines[index + 1]:
            continue                                  # a cue identifier (SRT number or VTT id)
        cue = re.sub(r"<v(?:\.[^\s>]+)?\s+([^>]+)>", r"\1: ", line)
        cue = html.unescape(re.sub(r"</?[^>]+>", "", cue)).strip()
        if cue and (not out or out[-1] != cue):
            out.append(cue)
    return "\n".join(out)


def markdown_table(rows: Sequence[Sequence[Any]]) -> Tuple[str, bool]:
    """A Markdown table, first row as header; empty rows and trailing empty columns dropped. Returns
    (table, truncated): rows past MAX_SHEET_ROWS and columns past the column cap are cut."""
    cells: List[List[str]] = []
    truncated = False
    for row in rows:
        values = [_cell(value) for value in row]
        if len(values) > _MAX_TABLE_COLUMNS:
            values, truncated = values[:_MAX_TABLE_COLUMNS], True
        if any(values):
            cells.append(values)
    if len(cells) > MAX_SHEET_ROWS:
        cells, truncated = cells[:MAX_SHEET_ROWS], True
    if not cells:
        return "", truncated
    width = max((max((i + 1 for i, v in enumerate(row) if v), default=0) for row in cells), default=0)
    padded = [(row + [""] * width)[:width] for row in cells]
    lines = ["| " + " | ".join(padded[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in padded[1:]]
    return "\n".join(lines), truncated


# -- readers --------------------------------------------------------------------------------------------------------

def _read_markitdown(path: str, kind: str, *, charset: Optional[str] = None) -> str:
    with open(path, "rb") as fh:
        return _convert(fh, kind, path, charset=charset)


def _convert(stream: BinaryIO, kind: str, local_path: Optional[str], *, charset: Optional[str] = None) -> str:
    from markitdown import StreamInfo
    from markitdown import converters

    class_name, extension, mimetype = MARKITDOWN_CONVERTERS[kind]
    converter = getattr(converters, class_name)()
    info = StreamInfo(extension=extension, mimetype=mimetype, charset=charset, local_path=local_path,
                      filename=os.path.basename(local_path) if local_path else None)
    return str(converter.convert(stream, info).markdown or "")


def _read_pdf(path: str, kind: str) -> _Read:
    from pypdf import PdfReader

    pages_total = 0
    try:
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise _Encrypted("user password")
        pages_total = _page_count(reader)
    except _Encrypted:
        raise
    except Exception as exc:  # pypdf cannot count a damaged file pdfminer may still read; density then is unknown
        if is_memory_failure(exc):
            raise
        pages_total = 0
    text, truncated = _pdfminer_text(path)
    return text, pages_total, truncated


def _pdfminer_text(path: str) -> Tuple[str, bool]:
    """pdfminer's extract_text, page by page: layout stops once MAX_TEXT_CHARS characters are out (R2 P-m6), so a
    thousand dense pages never all go through layout for text the cap drops. `truncated` when a page was left."""
    from pdfminer.converter import TextConverter
    from pdfminer.layout import LAParams
    from pdfminer.pdfinterp import PDFPageInterpreter, PDFResourceManager
    from pdfminer.pdfpage import PDFPage

    with open(path, "rb") as fh, io.StringIO() as output:
        manager = PDFResourceManager(caching=True)
        device = TextConverter(manager, output, codec="utf-8", laparams=LAParams())
        interpreter = PDFPageInterpreter(manager, device)
        pages = PDFPage.get_pages(fh, caching=True)
        for page in pages:
            interpreter.process_page(page)
            if output.tell() >= MAX_TEXT_CHARS:      # StringIO's position: the characters written so far
                return output.getvalue(), next(pages, None) is not None
        return output.getvalue(), False


def _page_count(reader: Any) -> int:
    """Pages found by walking the page tree. pypdf answers `len(reader.pages)` of an encrypted file from the root's
    `/Count`, which a file can set to anything (negative, a billion); the walk is bounded by pypdf's page-tree
    limits."""
    if reader.flattened_pages is None:
        try:
            reader.get_page(0)                       # flattens the tree (pypdf's public way in)
        except IndexError:
            return 0                                 # the tree holds no page
    return max(len(reader.flattened_pages or ()), 0)


def _read_converter(path: str, kind: str) -> _Read:
    return _read_markitdown(path, kind), 0, False


def _read_epub_package(path: str) -> _Read:
    """An Apple Books package (§5.1): its tar repacked as an EPUB ZIP in memory, then the EPUB path, DRM check
    included. The ZIP caps were applied to the tar's headers by the repack."""
    package = epub_zip_from_tar(path)
    with zipfile.ZipFile(package) as archive:
        if _epub_encrypted(archive, archive.namelist()):
            raise _Encrypted("EPUB DRM")
    package.seek(0)
    return _convert(package, "epub", None), 0, False


def _read_html(path: str, kind: str) -> _PageRead:
    data, cut = _read_prefix(path)
    text, script_rendered = html_page_text(decode_text(data, cut))
    return _PageRead(text, 0, cut, script_rendered)


def _read_text(path: str, kind: str) -> Union[_Read, _PageRead]:
    """Text as decoded; a Markdown or text file that is HTML (`text_is_html`) goes through the HTML converter whole
    (§25.3, parser@2). If that path fails, the file keeps its decoded text, as v1 stored it (review m3); a memory
    failure still ends the child with its memory code."""
    data, cut = _read_prefix(path)
    text = decode_text(data, cut)
    if kind in _HTML_CAPABLE_TEXT_KINDS and text_is_html(text):
        try:
            page, script_rendered = html_page_text(text)
        except Exception as exc:  # the HTML reading is an improvement over the text, never a new way to fail
            if is_memory_failure(exc):
                raise
            return text, 0, cut
        return _PageRead(page, 0, cut, script_rendered)
    return text, 0, cut


def _read_subtitles(path: str, kind: str) -> _Read:
    data, cut = _read_prefix(path)
    return subtitle_cues(decode_text(data, cut)), 0, cut


def _read_csv(path: str, kind: str) -> _Read:
    data, cut = _read_prefix(path)
    text = decode_text(data, cut)
    try:
        dialect: Any = csv.Sniffer().sniff(text[:_CSV_SNIFF_CHARS], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: List[List[str]] = []
    try:
        for row in csv.reader(io.StringIO(text), dialect):
            rows.append(row)
            if len(rows) > MAX_SHEET_ROWS:
                break
    except csv.Error:
        return text, 0, cut                           # not parseable as a table: indexed as text
    table, truncated = markdown_table(rows)
    return table, 0, cut or truncated


def _read_rtf(path: str, kind: str) -> _Read:
    from striprtf.striprtf import rtf_to_text

    data, cut = _read_prefix(path)
    return str(rtf_to_text(data.decode("latin-1"), errors="replace")), 0, cut


def _read_eml(path: str, kind: str) -> _Read:
    import email
    import email.policy
    from email.message import EmailMessage

    with open(path, "rb") as fh:
        message = email.message_from_binary_file(fh, policy=email.policy.default)
    if not isinstance(message, EmailMessage):
        raise ValueError("not a MIME message")
    header = [f"{name}: {message[name]}" for name in ("Subject", "From", "To", "Cc", "Date") if message[name]]
    body = message.get_body(preferencelist=("plain", "html"))
    text = ""
    if isinstance(body, EmailMessage):
        content = _part_text(body)
        if body.get_content_subtype() == "html":
            from markitdown.converters import HtmlConverter

            content = str(HtmlConverter().convert_string(content).markdown or "")
        text = content
    return "\n".join(header + ["", text]).strip(), 0, False


def _read_xlsx(path: str, kind: str) -> _Read:
    """Rows are read at most _MAX_TABLE_COLUMNS + 1 wide (a cell in column XFD must not pad every row to 16384
    values), each sheet is rendered as soon as it is read, and reading stops once MAX_TEXT_CHARS characters are
    collected: what the parser holds stays near the text cap whatever the sheet declares. A value past the read bound
    is content the text does not hold: the outcome is `truncated` (arbiter R1-3). The declared dimension is not read
    for that, nor is a formatted-empty cell: a sheet read in full is not `truncated` (R2 P-m3)."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks: List[str] = []
    collected = 0
    truncated = False
    try:
        for sheet in book.worksheets:
            if collected >= MAX_TEXT_CHARS:
                truncated = True
                break
            bound = _MAX_TABLE_COLUMNS + 1
            sheet.reset_dimensions()                  # a declared A1:XFD1048576 must not pad every row
            past_bound = _watch_far_cells(sheet, bound)
            rows: List[Sequence[Any]] = []
            missing: Optional[Sequence[Any]] = None
            for row in sheet.iter_rows(max_col=_MAX_TABLE_COLUMNS + 1, values_only=True):
                if row is missing:
                    continue                          # openpyxl yields one shared tuple for every absent row
                if row.count(None) == len(row):
                    missing = row
                    continue
                values = _trim_row(row)
                texts = [str(value) for value in values if value is not None]
                if not any(text.strip() for text in texts):
                    continue
                rows.append(values)
                collected += sum(len(text) + 3 for text in texts)     # "| " separators included
                if len(rows) > MAX_SHEET_ROWS:
                    break
                if collected >= MAX_TEXT_CHARS:
                    truncated = True
                    break
            table, cut = markdown_table(rows)
            truncated = truncated or cut or past_bound[0]
            if table:
                blocks.append(f"## {sheet.title}\n\n{table}")
    finally:
        book.close()
    return "\n\n".join(blocks), 0, truncated


def _watch_far_cells(sheet: Any, bound: int) -> List[bool]:
    """A flag set once any row of `sheet` holds a value past column `bound` (a cell that only carries a style, or
    blank text, is no content). iter_rows(max_col=bound) drops those cells without a trace, so the sheet's row
    builder is wrapped: it receives every parsed cell of a row (already in memory; nothing is padded) before cutting
    it to `bound`. openpyxl's `_get_row` is private, so openpyxl is pinned
    (requirements.txt) and tests/unit/adapters/test_drive_parser_xlsx_far_cell_hook.py fails if the hook stops
    seeing cells."""
    seen = [False]
    build_row = sheet._get_row

    def _get_row(row: Sequence[Mapping[str, Any]], min_col: int = 1, max_col: Optional[int] = None,
                 values_only: bool = False) -> Any:
        if not seen[0] and any(int(cell.get("column") or 0) > bound and _holds_value(cell.get("value"))
                               for cell in row):
            seen[0] = True
        return build_row(row, min_col, max_col, values_only)

    sheet._get_row = _get_row
    return seen


def _holds_value(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _trim_row(row: Sequence[Any]) -> Sequence[Any]:
    """The row without its trailing empty cells (iter_rows pads every row to max_col)."""
    end = len(row)
    while end and row[end - 1] is None:
        end -= 1
    return row[:end]


def _read_xls(path: str, kind: str) -> _Read:
    import xlrd

    try:
        book = xlrd.open_workbook(path, on_demand=True)
    except xlrd.XLRDError as exc:
        if "encrypted" in str(exc).lower():
            raise _Encrypted(str(exc)) from exc
        raise
    try:
        sheets: List[Tuple[str, List[Sequence[Any]], bool]] = []
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            rows: List[Sequence[Any]] = []
            for r in range(min(sheet.nrows, MAX_SHEET_ROWS + 1)):
                rows.append([_xls_value(cell, book.datemode) for cell in sheet.row(r)])
            sheets.append((sheet.name, rows, sheet.nrows > MAX_SHEET_ROWS + 1))
            book.unload_sheet(index)
    finally:
        book.release_resources()
    return _sheets_text(sheets)


def _read_odf(path: str, kind: str) -> _Read:
    from defusedxml import ElementTree

    with zipfile.ZipFile(path) as archive:
        content = archive.read("content.xml")
    root = ElementTree.fromstring(content)
    body = root.find(f"{{{_ODF['office']}}}body")
    if body is None:
        return "", 0, False
    blocks: List[str] = []
    truncated = _odf_blocks(body, kind, blocks)
    return "\n\n".join(block for block in blocks if block.strip()), 0, truncated


_READERS: Dict[str, Callable[[str, str], Union[_Read, _PageRead]]] = {
    "pdf": _read_pdf, "docx": _read_converter, "pptx": _read_converter, "epub": _read_converter,
    "msg": _read_converter, "xlsx": _read_xlsx, "xls": _read_xls, "odt": _read_odf, "ods": _read_odf,
    "odp": _read_odf, "html": _read_html, "text": _read_text, "markdown": _read_text, "json": _read_text,
    "csv": _read_csv, "subtitles": _read_subtitles, "rtf": _read_rtf, "eml": _read_eml,
}


# -- internals ------------------------------------------------------------------------------------------------------

def _check_container(path: str, kind: str) -> None:
    """ZIP caps (Phase 0) and encryption, before any parser opens the package."""
    with open(path, "rb") as fh:
        magic = fh.read(len(_OLE_MAGIC))
    if magic == _OLE_MAGIC:
        if kind in _OOXML_KINDS and "EncryptedPackage" in _ole_streams(path):
            raise _Encrypted("encrypted OOXML package")
        raise _WrongContainer("an OLE2 file under a ZIP kind")
    check_zip_container(path)
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if kind == "epub" and _epub_encrypted(archive, names):
            raise _Encrypted("EPUB DRM")
        if kind in ("odt", "ods", "odp") and _odf_encrypted(archive, names):
            raise _Encrypted("ODF password")


def _deterministic(exc: Exception) -> bool:
    """A failure the same bytes always produce: `unsupported` (terminal, §8.2), not `error` (retried MAX_ATTEMPTS
    times). Nesting too deep for the parser, a container that is not one, XML with a DTD or entities."""
    from defusedxml import DefusedXmlException

    return isinstance(exc, (RecursionError, zipfile.BadZipFile, _WrongContainer, PackageError, DefusedXmlException))


def _ole_streams(path: str) -> FrozenSet[str]:
    import olefile

    if not olefile.isOleFile(path):
        return frozenset()
    with olefile.OleFileIO(path) as ole:
        return frozenset("/".join(entry) for entry in ole.listdir())


def _epub_encrypted(archive: zipfile.ZipFile, names: Iterable[str]) -> bool:
    names = set(names)
    if "META-INF/rights.xml" in names:
        return True
    if "META-INF/encryption.xml" not in names:
        return False
    from defusedxml import ElementTree

    root = ElementTree.fromstring(archive.read("META-INF/encryption.xml"))
    uris = [el.get("URI", "") for el in root.iter() if el.tag.endswith("CipherReference")]
    return any(not uri.lower().endswith(_FONT_SUFFIXES) for uri in uris)


def _odf_encrypted(archive: zipfile.ZipFile, names: Iterable[str]) -> bool:
    if "META-INF/manifest.xml" not in set(names):
        return False
    from defusedxml import ElementTree

    root = ElementTree.fromstring(archive.read("META-INF/manifest.xml"))
    return any(el.tag.endswith("}encryption-data") for el in root.iter())


def _odf_blocks(node: Any, kind: str, out: List[str]) -> bool:
    """Paragraphs, headings, tables and slides in document order; returns whether a table was cut."""
    truncated = False
    for child in node:
        tag = child.tag
        if tag in (_q("text", "p"), _q("text", "h")):
            line = _odf_inline(child).strip()
            if tag == _q("text", "h") and line:
                level = int(child.get(_q("text", "outline-level"), "1") or "1")
                line = "#" * max(1, min(level, 6)) + " " + line
            out.append(line)
        elif tag == _q("table", "table"):
            table, cut = markdown_table(_odf_rows(child))
            truncated = truncated or cut
            if kind == "ods":
                out.append(f"## {child.get(_q('table', 'name'), '')}".rstrip())
            out.append(table)
        elif tag == _q("draw", "page"):
            out.append(f"## {child.get(_q('draw', 'name'), '')}".rstrip())
            truncated = _odf_blocks(child, kind, out) or truncated
        else:
            truncated = _odf_blocks(child, kind, out) or truncated
    return truncated


def _odf_rows(table: Any) -> List[List[str]]:
    rows: List[List[str]] = []
    for row in table.iter(_q("table", "table-row")):
        cells: List[str] = []
        pending_empty = 0
        for cell in row:
            if cell.tag not in (_q("table", "table-cell"), _q("table", "covered-table-cell")):
                continue
            repeat = min(int(cell.get(_q("table", "number-columns-repeated"), "1") or "1"), _MAX_TABLE_COLUMNS + 1)
            value = " ".join(_odf_inline(p).strip() for p in cell.iter(_q("text", "p"))).strip()
            if not value:
                pending_empty += repeat              # trailing empties are never materialised
                continue
            cells.extend([""] * min(pending_empty, _MAX_TABLE_COLUMNS + 1))
            pending_empty = 0
            cells.extend([value] * repeat)
            if len(cells) > _MAX_TABLE_COLUMNS:
                break
        if not any(cells):
            continue                                 # empty rows, however often repeated, cost nothing
        repeat = min(int(row.get(_q("table", "number-rows-repeated"), "1") or "1"), MAX_SHEET_ROWS + 1)
        rows.extend([cells] * repeat)
        if len(rows) > MAX_SHEET_ROWS:
            break
    return rows


def _odf_inline(node: Any) -> str:
    parts = [node.text or ""]
    for child in node:
        if child.tag == _q("text", "s"):
            parts.append(" " * min(int(child.get(_q("text", "c"), "1") or "1"), 100))
        elif child.tag == _q("text", "tab"):
            parts.append("\t")
        elif child.tag == _q("text", "line-break"):
            parts.append("\n")
        elif child.tag != _q("text", "note"):
            parts.append(_odf_inline(child))
        parts.append(child.tail or "")
    return "".join(parts)


def _q(namespace: str, name: str) -> str:
    return f"{{{_ODF[namespace]}}}{name}"


def _sheets_text(sheets: Sequence[Tuple[str, List[Sequence[Any]], bool]]) -> _Read:
    blocks: List[str] = []
    truncated = False
    for title, rows, cut in sheets:
        table, table_cut = markdown_table(rows)
        truncated = truncated or cut or table_cut
        if table:
            blocks.append(f"## {title}\n\n{table}")
    return "\n\n".join(blocks), 0, truncated


def _xls_value(cell: Any, datemode: int) -> Any:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except (ValueError, OverflowError):
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return " ".join(str(value).split()).replace("|", "\\|")


def _part_text(part: Any) -> str:
    try:
        return str(part.get_content())
    except (LookupError, UnicodeError, AttributeError, KeyError):
        payload = part.get_payload(decode=True) or b""
        return decode_text(payload) if isinstance(payload, bytes) else str(payload)


def _read_prefix(path: str) -> Tuple[bytes, bool]:
    """At most what MAX_TEXT_CHARS characters can take (4 bytes each); `cut` when the file is longer."""
    limit = MAX_TEXT_CHARS * 4
    with open(path, "rb") as fh:
        data = fh.read(limit + 1)
    return (data[:limit], True) if len(data) > limit else (data, False)


def _guess_encoding(data: bytes) -> str:
    """charset-normalizer's ranking, except that a common legacy code page whose chaos is within _CHAOS_MARGIN of
    the best wins: on a short file the detector ties CP1251 with mac_cyrillic or a CJK code page."""
    from charset_normalizer import from_bytes

    matches = list(from_bytes(data))
    if not matches:
        return "utf-8"
    best = matches[0]
    near = [m for m in matches if m.encoding in _PREFERRED_LEGACY and m.chaos <= best.chaos + _CHAOS_MARGIN]
    if not near:
        return str(best.encoding)
    sample = data[:_PLAUSIBILITY_SAMPLE_BYTES]
    return str(max(near, key=lambda m: (-_implausible_words(sample.decode(m.encoding, errors="replace")),
                                        m.coherence, -_PREFERRED_LEGACY.index(m.encoding))).encoding)


def _implausible_words(text: str) -> int:
    """Words no language writes: Latin and Cyrillic letters in one word ("Rйsumй": CP1252 read as CP1251), or a
    lowercase letter followed by a capital ("пРМБФБ": KOI8 read as CP1251, which swaps the case). ASCII decodes
    the same in every candidate, so a real "iPhone" costs all of them alike."""
    count = 0
    latin = cyrillic = inverted = False
    previous_lower = False
    for char in text + " ":
        if not char.isalpha():
            count += (latin and cyrillic) or inverted
            latin = cyrillic = inverted = previous_lower = False
            continue
        if "\u0400" <= char <= "\u04ff":
            cyrillic = True
        elif char < "\u0250":
            latin = True
        inverted = inverted or (previous_lower and char.isupper())
        previous_lower = char.islower()
    return count


def _decode(data: bytes, encoding: str, cut: bool, errors: str) -> str:
    decoder = codecs.getincrementaldecoder(encoding)(errors=errors)
    return str(decoder.decode(data, final=not cut))


def _finish(text: str, pages_total: int, truncated: bool, *, script_rendered: bool = False) -> ExtractionOutcome:
    text = text.replace("\x00", "")
    text = "\n".join(line.rstrip() for line in re.split(r"\r?\n", text))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_TEXT_CHARS:
        text, truncated = text[:MAX_TEXT_CHARS], True
    if not text:
        return ExtractionOutcome(status=ExtractionStatus.NO_TEXT, pages_total=pages_total, truncated=truncated)
    return ExtractionOutcome(status=ExtractionStatus.TEXT, text=text, pages_total=pages_total, truncated=truncated,
                             script_rendered=script_rendered)
