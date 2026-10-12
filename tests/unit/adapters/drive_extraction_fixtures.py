"""Small documents generated at test time for the drive parser (USER_DRIVE_INDEXING_RFC §5.1, §9): no binary blobs
in the repository. Each builder writes one file and returns its path."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Iterable, Optional, Sequence

from src.utils.document_hardening import MAX_ZIP_MEMBERS, MAX_ZIP_UNCOMPRESSED_BYTES


def make_pdf(path: Path, pages: Sequence[str]) -> Path:
    """A valid PDF with one Helvetica text line per page (empty string: a page with no text layer)."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1")
        stream = b"BT /F1 12 Tf 72 720 Td (" + escaped + b") Tj ET" if text else b""
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_id = len(objects)
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >>"
                       b" /Contents %d 0 R >>" % content_id)
        kids.append(b"%d 0 R" % len(objects))
    objects[1] = b"<< /Type /Pages /Kids [" + b" ".join(kids) + b"] /Count %d >>" % len(pages)
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))
    return path


def encrypt_pdf(source: Path, target: Path, *, user_password: str, owner_password: str = "owner") -> Path:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(str(source)))
    writer.encrypt(user_password=user_password, owner_password=owner_password, algorithm="RC4-128")
    with open(target, "wb") as fh:
        writer.write(fh)
    return target


def make_docx(path: Path, paragraphs: Iterable[str]) -> Path:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(str(path))
    return path


def make_pptx(path: Path, titles: Iterable[str]) -> Path:
    import pptx

    deck = pptx.Presentation()
    for title in titles:
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = title
    deck.save(str(path))
    return path


def make_xlsx(path: Path, sheets: Sequence[tuple]) -> Path:
    """`sheets`: (title, rows) pairs."""
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    book.save(str(path))
    return path


def make_epub(path: Path, body_html: str, *, encryption_xml: Optional[str] = None,
              rights: bool = False) -> Path:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>')
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
                   'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
                   '<dc:title>Test Book</dc:title><dc:identifier id="id">x</dc:identifier></metadata>'
                   '<manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/c1.xhtml",
                   f'<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><body>{body_html}</body></html>')
        if encryption_xml is not None:
            z.writestr("META-INF/encryption.xml", encryption_xml)
        if rights:
            z.writestr("META-INF/rights.xml", "<rights/>")
    return path


ODF_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"')
_ODF_MIME = {"odt": "text", "ods": "spreadsheet", "odp": "presentation"}


def make_odf(path: Path, kind: str, body: str, *, encrypted: bool = False) -> Path:
    """An ODF package whose office:body holds `body` (already namespaced XML)."""
    manifest_extra = ('<manifest:encryption-data manifest:checksum-type="SHA1" manifest:checksum="x"/>'
                      if encrypted else "")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), f"application/vnd.oasis.opendocument.{_ODF_MIME[kind]}",
                   compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", f'<?xml version="1.0"?><office:document-content {ODF_NS}>'
                                  f'<office:body>{body}</office:body></office:document-content>')
        z.writestr("META-INF/manifest.xml",
                   '<?xml version="1.0"?><manifest:manifest '
                   'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
                   '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml">'
                   f'{manifest_extra}</manifest:file-entry></manifest:manifest>')
    return path


def make_member_bomb(path: Path, prefix: str = "") -> Path:
    """Phase 0's member-count bomb: more ZIP members than MAX_ZIP_MEMBERS (`prefix="word/"`: sniffed as a DOCX)."""
    with zipfile.ZipFile(path, "w") as z:
        for i in range(MAX_ZIP_MEMBERS + 1):
            z.writestr(f"{prefix}m{i}", b"")
    return path


def make_inflation_bomb(path: Path, member: str = "word/document.xml") -> Path:
    """Phase 0's size bomb: a small archive declaring more than MAX_ZIP_UNCOMPRESSED_BYTES inflated."""
    chunk = b"\0" * (1024 * 1024)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        with z.open(member, "w", force_zip64=True) as out:
            for _ in range(MAX_ZIP_UNCOMPRESSED_BYTES // len(chunk) + 1):
                out.write(chunk)
    return path
