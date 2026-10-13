"""An xlsx is `truncated` only when a cell past the column read bound holds a value (round-2 review P-m3; arbiter:
the declared-dimension signal is dropped). A formatted-empty cell in column 2000, or a bogus declared A1:XFD1048576,
made search tell the model "only the beginning was indexed" for a file read in full."""
import re
import zipfile

from src.adapters.drive_parser.readers import read_content
from src.domain.drive_extraction import ExtractionStatus


def _book(path, cells, styled=()):
    import openpyxl
    from openpyxl.styles import PatternFill

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet"
    for row, column, value in cells:
        sheet.cell(row=row, column=column, value=value)
    for row, column in styled:
        sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor="FFFF00")
    book.save(str(path))
    return path


def _declare(path, ref, target):
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w") as out:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text, count = re.subn(r'<dimension ref="[^"]*"', f'<dimension ref="{ref}"', data.decode())
                assert count == 1
                data = text.encode()
            out.writestr(item, data)
    return target


def test_a_formatted_empty_cell_in_column_2000_is_not_truncation(tmp_path):
    path = _book(tmp_path / "s.xlsx", [(1, 1, "Item"), (2, 1, "Rent")], styled=[(1, 2000)])
    outcome = read_content(str(path), "xlsx")
    assert outcome.status is ExtractionStatus.TEXT
    assert "| Rent |" in outcome.text
    assert outcome.truncated is False


def test_a_bogus_declared_whole_sheet_dimension_is_not_truncation(tmp_path):
    plain = _book(tmp_path / "p.xlsx", [(1, 1, "Item"), (1, 2, "Cost"), (2, 1, "Rent"), (2, 2, 1200)])
    declared = _declare(plain, "A1:XFD1048576", tmp_path / "declared.xlsx")
    outcome = read_content(str(declared), "xlsx")
    assert "| Rent | 1200 |" in outcome.text
    assert outcome.truncated is False


def test_a_blank_string_past_the_bound_is_not_truncation(tmp_path):
    path = _book(tmp_path / "s.xlsx", [(1, 1, "Item"), (2, 1, "Rent"), (2, 2000, "   ")])
    assert read_content(str(path), "xlsx").truncated is False


def test_a_value_past_the_bound_beside_a_formatted_empty_one_still_marks_truncated(tmp_path):
    path = _book(tmp_path / "s.xlsx", [(1, 1, "Item"), (2, 1, "Rent"), (3, 1800, "hidden")], styled=[(1, 2000)])
    assert read_content(str(path), "xlsx").truncated is True
