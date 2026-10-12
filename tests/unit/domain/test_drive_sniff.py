"""Kind from the downloaded head (USER_DRIVE_INDEXING_RFC §5.1, §5.5 class 2, §8.1 step 4)."""
import io
import struct
import zipfile

import pytest

from src.domain.drive_sniff import SNIFF_HEAD_BYTES, sniff_kind


def _head(data: bytes) -> bytes:
    return data[:SNIFF_HEAD_BYTES]


def _zip(members, *, mimetype: bytes = b"") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        if mimetype:
            z.writestr(zipfile.ZipInfo("mimetype"), mimetype, compress_type=zipfile.ZIP_STORED)
        for name, data in members:
            z.writestr(name, data)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    import docx
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("hello")
    document.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl
    buf = io.BytesIO()
    book = openpyxl.Workbook()
    book.active["A1"] = "x"
    book.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    import pptx
    buf = io.BytesIO()
    pptx.Presentation().save(buf)
    return buf.getvalue()


_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _ole(*names: str) -> bytes:
    body = b"".join(n.encode("utf-16-le") + b"\x00" * 8 for n in names)
    return _OLE + b"\x00" * 504 + body


# =============================================================================
# PDF, RTF, images and other binaries: magic bytes win over the name
# =============================================================================

@pytest.mark.parametrize("declared", ["pdf", "text", "docx", "other"])
def test_pdf_magic_wins(declared):
    assert sniff_kind(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n1 0 obj", declared) == "pdf"


def test_pdf_header_after_leading_junk():
    assert sniff_kind(b"\x00" * 500 + b"%PDF-1.4\n", "pdf") == "pdf"


def test_pdf_header_past_one_kib_is_not_a_pdf():
    assert sniff_kind(b"\x00" * 1100 + b"%PDF-1.4\n", "pdf") == "other"


def test_declared_pdf_that_is_text_is_other():
    assert sniff_kind(b"just some words", "pdf") == "other"


def test_rtf():
    assert sniff_kind(b"{\\rtf1\\ansi\\ansicpg1251 hello}", "rtf") == "rtf"
    assert sniff_kind(b"{\\rtf1\\ansi hello}", "text") == "rtf"


def test_declared_rtf_without_rtf_header_is_other():
    assert sniff_kind(b"plain words", "rtf") == "other"


@pytest.mark.parametrize("head,kind", [
    (b"\x89PNG\r\n\x1a\n\x00\x00", "image"),
    (b"\xff\xd8\xff\xe0\x00\x10JFIF", "image"),
    (b"GIF89a\x01\x00", "image"),
    (b"RIFF\x00\x00\x00\x00WEBPVP8 ", "image"),
    (b"\x00\x00\x00\x18ftypheic\x00\x00", "image"),
    (b"\x00\x00\x00\x18ftypisom\x00\x00", "video"),
    (b"\x00\x00\x00\x18ftypM4A \x00\x00", "audio"),
    (b"RIFF\x00\x00\x00\x00WAVEfmt ", "audio"),
    (b"\x1f\x8b\x08\x00", "archive"),
    (b"7z\xbc\xaf\x27\x1c\x00\x04", "archive"),
    (b"Rar!\x1a\x07\x00", "archive"),
    (b"\x7fELF\x02\x01", "executable"),
    (b"MZ\x90\x00\x03\x00", "executable"),
    (b"BZh91AY&SY\x00\x01", "archive"),
    (b"ID3\x04\x00\x00", "audio"),
    (b"SQLite format 3\x00", "database"),
    (b"\x1aE\xdf\xa3\x01\x00", "video"),
])
def test_binary_signatures(head, kind):
    assert sniff_kind(head, "text") == kind


def test_tar_magic_at_257():
    head = b"a.txt".ljust(257, b"\x00") + b"ustar\x0000"
    assert sniff_kind(head, "text") == "archive"


@pytest.mark.parametrize("text", [b"MZ notes from the meeting", b"BZh is a note", b"ID3 tags explained"])
def test_printable_magic_in_a_text_head_stays_text(text):
    assert sniff_kind(text, "text") == "text"


def test_declared_image_without_known_magic_stays_image():
    assert sniff_kind(b"<svg xmlns='http://www.w3.org/2000/svg'></svg>", "image") == "image"


# =============================================================================
# ZIP containers: OOXML, ODF, EPUB told apart by local headers in the head
# =============================================================================

def test_txt_that_is_a_zip_is_an_archive():
    assert sniff_kind(_head(_zip([("a.txt", b"hello")])), "text") == "archive"


def test_real_docx_xlsx_pptx():
    assert sniff_kind(_head(_docx_bytes()), "docx") == "docx"
    assert sniff_kind(_head(_xlsx_bytes()), "xlsx") == "xlsx"
    assert sniff_kind(_head(_pptx_bytes()), "pptx") == "pptx"


def test_ooxml_part_names_beat_the_declared_kind():
    data = _zip([("[Content_Types].xml", b"<Types/>"), ("xl/workbook.xml", b"<workbook/>")])
    assert sniff_kind(data, "docx") == "xlsx"


def test_ooxml_with_only_content_types_visible_keeps_declared_ooxml_kind():
    data = _zip([("[Content_Types].xml", b"<Types/>" + b"x" * 64)])
    assert sniff_kind(data, "pptx") == "pptx"


def test_ooxml_with_only_content_types_visible_declared_text_is_archive():
    data = _zip([("[Content_Types].xml", b"<Types/>")])
    assert sniff_kind(data, "text") == "archive"


@pytest.mark.parametrize("mime,kind", [
    (b"application/epub+zip", "epub"),
    (b"application/vnd.oasis.opendocument.text", "odt"),
    (b"application/vnd.oasis.opendocument.spreadsheet", "ods"),
    (b"application/vnd.oasis.opendocument.presentation", "odp"),
    (b"application/vnd.oasis.opendocument.graphics", "archive"),
])
def test_odf_and_epub_by_stored_mimetype(mime, kind):
    data = _zip([("content.xml", b"<x/>")], mimetype=mime)
    assert sniff_kind(data, "docx") == kind


def test_zip_head_too_short_to_name_a_member_keeps_a_zip_kind():
    assert sniff_kind(b"PK\x03\x04\x14\x00", "epub") == "epub"
    assert sniff_kind(b"PK\x03\x04\x14\x00", "text") == "archive"


def test_empty_zip():
    assert sniff_kind(b"PK\x05\x06" + b"\x00" * 18, "odt") == "odt"


def test_data_descriptor_stops_the_walk():
    name = b"word/document.xml"
    header = b"PK\x03\x04" + struct.pack("<HHHHHIIIHH", 20, 0x08, 8, 0, 0, 0, 0, 0, len(name), 0) + name
    assert sniff_kind(header + b"\x00" * 40, "xlsx") == "docx"


def test_iwork_zip_keeps_iwork():
    assert sniff_kind(_zip([("Index/Document.iwa", b"\x00")]), "iwork") == "iwork"


# =============================================================================
# OLE2: xls / msg / legacy Office / encrypted OOXML
# =============================================================================

@pytest.mark.parametrize("names,declared,kind", [
    (("Workbook",), "text", "xls"),
    (("__substg1.0_0037001F",), "other", "msg"),
    (("WordDocument",), "docx", "doc"),
    (("PowerPoint Document",), "pptx", "ppt"),
    (("EncryptedPackage", "EncryptionInfo"), "xlsx", "xlsx"),
    (("EncryptedPackage",), "text", "other"),
])
def test_ole_directory_names_in_the_head(names, declared, kind):
    assert sniff_kind(_ole(*names), declared) == kind


@pytest.mark.parametrize("declared,kind", [
    ("xls", "xls"), ("msg", "msg"), ("docx", "docx"), ("pdf", "other"), ("text", "other"),
])
def test_ole_without_visible_names_falls_back_to_a_compatible_declared_kind(declared, kind):
    assert sniff_kind(_OLE + b"\x00" * 600, declared) == kind


# =============================================================================
# Text kinds
# =============================================================================

@pytest.mark.parametrize("declared", ["text", "markdown", "csv", "json", "subtitles", "eml", "html"])
def test_text_heads_keep_their_declared_kind(declared):
    assert sniff_kind("Привіт, світ\nline two".encode("utf-8"), declared) == declared


@pytest.mark.parametrize("bom", [b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff"])
def test_bom_text_is_text_even_with_nul_bytes(bom):
    assert sniff_kind(bom + "hello".encode("utf-16-le"), "text") == "text"


def test_cp1251_text_is_text():
    assert sniff_kind("Договір оренди".encode("cp1251"), "text") == "text"


def test_binary_head_declared_text_is_other():
    assert sniff_kind(b"\x00\x01\x02\x03binary", "text") == "other"


@pytest.mark.parametrize("head", [b"<!DOCTYPE html><html>", b"\xef\xbb\xbf\n  <HTML lang='en'>"])
def test_text_file_holding_html_is_converted_as_html(head):
    assert sniff_kind(head, "text") == "html"
    assert sniff_kind(head, "markdown") == "html"


def test_html_detection_does_not_touch_other_text_kinds():
    assert sniff_kind(b"<!DOCTYPE html>", "json") == "json"


@pytest.mark.parametrize("declared", ["code", "log", "data_dump"])
def test_non_admitted_text_kinds_stay_themselves(declared):
    assert sniff_kind(b"select * from t;", declared) == declared


def test_unknown_declared_kind_is_other():
    assert sniff_kind(b"words", "other") == "other"
    assert sniff_kind(b"", "text") == "text"
